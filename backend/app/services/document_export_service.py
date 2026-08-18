"""Small, dependency-backed renderers for downloadable business reports."""

from __future__ import annotations

import csv
import io
import os
from datetime import date, datetime
from pathlib import Path
from typing import Mapping, Sequence

from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFError, TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


Column = tuple[str, str]
Section = tuple[str, Sequence[Column], Sequence[Mapping[str, object]]]

_FORMATS = {"csv", "xlsx", "pdf"}
_MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
_PDF_FONT_NAME: str | None = None
WINDOWS_CJK_FONT_CANDIDATES = ("simhei.ttf", "msyh.ttc", "msyhl.ttc", "simsun.ttc")


def download_response(title: str, sections: Sequence[Section], file_format: str, filename: str) -> StreamingResponse:
    """Render report sections and return a binary attachment response."""
    if file_format not in _FORMATS:
        raise HTTPException(status_code=400, detail="format must be one of: csv, xlsx, pdf")

    if file_format == "csv":
        payload = _render_csv(title, sections)
    elif file_format == "xlsx":
        payload = _render_xlsx(title, sections)
    else:
        payload = _render_pdf(title, sections)

    return StreamingResponse(
        io.BytesIO(payload),
        media_type=_MEDIA_TYPES[file_format],
        headers={"Content-Disposition": f'attachment; filename="{filename}.{file_format}"'},
    )


def _render_csv(title: str, sections: Sequence[Section]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([title])
    for section_title, columns, rows in sections:
        writer.writerow([])
        writer.writerow([section_title])
        writer.writerow([label for _, label in columns])
        for row in rows:
            writer.writerow([_safe_csv_value(row.get(key)) for key, _ in columns])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def _render_xlsx(title: str, sections: Sequence[Section]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "报表"
    worksheet.freeze_panes = "A3"
    worksheet.sheet_view.showGridLines = False

    current_row = 1
    max_columns = max((len(columns) for _, columns, _ in sections), default=1)
    worksheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_columns)
    title_cell = worksheet.cell(current_row, 1, title)
    title_cell.font = Font(bold=True, size=16, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="246B4B")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[current_row].height = 28
    current_row += 2

    for section_title, columns, rows in sections:
        worksheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(columns))
        section_cell = worksheet.cell(current_row, 1, section_title)
        section_cell.font = Font(bold=True, color="1B5E3C")
        section_cell.fill = PatternFill("solid", fgColor="E8F3EC")
        current_row += 1

        for column_index, (_, label) in enumerate(columns, start=1):
            cell = worksheet.cell(current_row, column_index, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3D8061")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

        for row in rows:
            for column_index, (key, _) in enumerate(columns, start=1):
                cell = worksheet.cell(current_row, column_index, _safe_spreadsheet_value(row.get(key)))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(row.get(key), (int, float)):
                    cell.number_format = "0.00"
            current_row += 1
        if not rows:
            worksheet.cell(current_row, 1, "暂无数据")
            current_row += 1
        current_row += 1

    for column_index in range(1, max_columns + 1):
        column_letter = get_column_letter(column_index)
        max_length = max(
            (len(str(cell.value or "")) for cell in worksheet[column_letter]),
            default=10,
        )
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 36)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _render_pdf(title: str, sections: Sequence[Section]) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
    )
    font_name = _pdf_font_name()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ExportTitle", parent=styles["Heading1"], fontName=font_name, fontSize=17, leading=22)
    section_style = ParagraphStyle("ExportSection", parent=styles["Heading2"], fontName=font_name, fontSize=11, leading=15, textColor=colors.HexColor("#1B5E3C"))
    table_style = ParagraphStyle("ExportCell", parent=styles["Normal"], fontName=font_name, fontSize=7.5, leading=10)
    story = [Paragraph(_escape_pdf(title), title_style), Spacer(1, 5 * mm)]

    for section_title, columns, rows in sections:
        story.append(Paragraph(_escape_pdf(section_title), section_style))
        table_rows = [[Paragraph(_escape_pdf(label), table_style) for _, label in columns]]
        table_rows.extend(
            [Paragraph(_escape_pdf(_display_value(row.get(key))), table_style) for key, _ in columns]
            for row in rows
        )
        if not rows:
            table_rows.append([Paragraph("暂无数据", table_style)] + [""] * (len(columns) - 1))

        available_width = landscape(A4)[0] - document.leftMargin - document.rightMargin
        column_widths = [available_width / max(len(columns), 1)] * len(columns)
        table = Table(table_rows, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3D8061")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7E4DB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend([table, Spacer(1, 4 * mm)])

    document.build(story)
    return buffer.getvalue()


def _pdf_font_name() -> str:
    global _PDF_FONT_NAME
    if _PDF_FONT_NAME:
        return _PDF_FONT_NAME

    for font_path in _windows_cjk_font_paths():
        font_name = _register_windows_cjk_font(font_path)
        if font_name:
            _PDF_FONT_NAME = font_name
            return _PDF_FONT_NAME

    _PDF_FONT_NAME = _register_cjk_cid_fallback()
    return _PDF_FONT_NAME


def _windows_cjk_font_paths() -> tuple[Path, ...]:
    windows_fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    return tuple(windows_fonts / filename for filename in WINDOWS_CJK_FONT_CANDIDATES)


def _register_windows_cjk_font(font_path: Path) -> str | None:
    if not font_path.is_file():
        return None

    font_name = f"SellHelpCJK_{font_path.stem}"
    try:
        pdfmetrics.registerFont(TTFont(font_name, str(font_path), subfontIndex=0))
    except (OSError, TTFError):
        return None
    return font_name


def _register_cjk_cid_fallback() -> str:
    font_name = "STSong-Light"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    return font_name


def _safe_csv_value(value: object) -> str:
    text = _display_value(value)
    leading_whitespace = text[: len(text) - len(text.lstrip())]
    unspaced = text[len(leading_whitespace):]
    return f"{leading_whitespace}'{unspaced}" if unspaced.startswith(("=", "+", "-", "@")) else text


def _safe_spreadsheet_value(value: object) -> object:
    text = _display_value(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return value if isinstance(value, (int, float)) else text


def _display_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float):
        return f"{value:.2f}"
    return str(value)


def _escape_pdf(value: object) -> str:
    return _display_value(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
