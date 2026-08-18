from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from reportlab.pdfbase import pdfmetrics

from app.models.all_models import Customer, Product, ProductBatch, SalesOrder, StockTake, WeeklyReport
from app.services import document_export_service as export_service


def seed_export_data(db):
    customer = Customer(name="导出客户", customer_type="小店", phone="13800000000")
    product = Product(name="导出商品", unit="箱", purchase_price=12.5)
    db.add_all([customer, product])
    db.flush()

    batch = ProductBatch(
        product_id=product.id,
        batch_no="EXP-001",
        purchase_price=12.5,
        total_quantity=10,
        remaining_quantity=8,
    )
    db.add(batch)
    db.flush()
    order = SalesOrder(
        order_no="SO-EXPORT-001",
        customer_id=customer.id,
        sale_date=datetime(2026, 8, 17, 9, 30),
        total_amount=100,
        final_amount=100,
        paid_amount=80,
        debt_amount=20,
        payment_type="部分结账",
        status="已完成",
    )
    stock_take = StockTake(
        take_date=date(2026, 8, 17),
        product_id=product.id,
        batch_id=batch.id,
        system_quantity=8,
        actual_quantity=7,
        diff_quantity=-1,
        diff_amount=-12.5,
        confirmed=True,
    )
    report = WeeklyReport(
        week_start=date(2026, 8, 10),
        week_end=date(2026, 8, 16),
        total_sales=100,
        total_profit=30,
        order_count=1,
        stock_value=100,
    )
    db.add_all([order, stock_take, report])
    db.commit()
    return customer, report


@pytest.mark.parametrize(
    ("path_template", "file_format", "content_type", "signature"),
    [
        ("/api/export/weekly-reports/{report_id}", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK"),
        ("/api/export/stock-report", "pdf", "application/pdf", b"%PDF-"),
        ("/api/export/customer-statement/{customer_id}", "csv", "text/csv", b"\xef\xbb\xbf"),
        ("/api/export/stock-takes", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK"),
        ("/api/export/sales-history", "pdf", "application/pdf", b"%PDF-"),
    ],
)
def test_report_exports_return_downloadable_binary_files(client, db_session, path_template, file_format, content_type, signature):
    customer, report = seed_export_data(db_session)
    path = path_template.format(customer_id=customer.id, report_id=report.id)

    response = client.get(path, params={"format": file_format})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(content_type)
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.content.startswith(signature)


def test_stock_report_json_remains_available_for_existing_integrations(client, db_session):
    seed_export_data(db_session)

    response = client.get("/api/export/stock-report")

    assert response.status_code == 200
    assert response.json()["total_products"] == 1
    assert response.json()["items"][0]["product_name"] == "导出商品"


def test_xlsx_export_opens_and_existing_history_and_config_apis_return_persisted_data(client, db_session):
    customer, report = seed_export_data(db_session)

    workbook_response = client.get(f"/api/export/weekly-reports/{report.id}", params={"format": "xlsx"})
    workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
    values = [cell for row in workbook.active.iter_rows(values_only=True) for cell in row if cell]
    stock_takes = client.get("/api/stock-takes")
    customer_sales = client.get(f"/api/customers/{customer.id}/sales")
    updated = client.put("/api/system/config/phase3_export_test", params={"value": "enabled", "description": "Phase 3 export verification"})
    config = client.get("/api/system/config")

    assert workbook_response.status_code == 200
    assert "周度经营报表" in values
    assert "经营概览" in values
    assert stock_takes.status_code == 200
    assert stock_takes.json()[0]["product_name"] == "导出商品"
    assert customer_sales.status_code == 200
    assert customer_sales.json()[0]["order_no"] == "SO-EXPORT-001"
    assert updated.status_code == 200
    assert config.status_code == 200
    assert config.json()["phase3_export_test"] == "enabled"


def test_pdf_font_falls_back_to_portable_cjk_font_when_windows_font_is_unavailable(monkeypatch):
    """Removing the CJK fallback would reintroduce unreadable Chinese PDFs outside a configured Windows font path."""
    monkeypatch.setattr(export_service, "WINDOWS_CJK_FONT_CANDIDATES", (), raising=False)
    monkeypatch.setattr(export_service, "_PDF_FONT_NAME", None)

    font_name = export_service._pdf_font_name()

    assert font_name == "STSong-Light"
    assert pdfmetrics.stringWidth("中文报表", font_name, 12) > 0
